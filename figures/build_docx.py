"""Build a combined manuscript+supplement DOCX from the two Markdown files using python-docx.
Handles headings, paragraphs, pipe tables, bold/italic/code, lists, page breaks, and
best-effort math rendering. Run from SSCR Submission/."""
import re, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIGMAP = {
    "**Figure 1.": "03_Figures/Figure1_framework.png",
    "**Figure 2.": "03_Figures/Figure2_recovery.png",
    "**Figure 3.": "03_Figures/Figure3_adversarial.png",
    "**Figure 4.": "03_Figures/Figure4_engagement.png",
    "**Figure S1.": "03_Figures/FigureS1_family_sizes.png",
    "**Figure S3.": "03_Figures/FigureS3_sensitivity.png",
    "**Figure S4.": "03_Figures/FigureS4_recovery_heatmap.png",
    "**Figure S5.": "03_Figures/FigureS5_paired_size.png",
    "**Figure S6.": "03_Figures/FigureS6_adversarial.png",
    "**Figure S9.": "03_Figures/FigureS9_attrition.png",
    "**Figure S10.": "03_Figures/FigureS10_sensitivity.png",
    "**Figure S11.": "03_Figures/FigureS11_decomposition_nulls.png",
    "**Figure S12.": "03_Figures/FigureS12_prospective.png",
}

MAN = "01_Manuscript/SSCR_Manuscript_Anonymized.md"
SUP = "02_Supplement/SSCR_Supplement.md"
OUT = "SSCR_Manuscript_and_Supplement.docx"

def demath(t):
    t = t.replace("\\(", "").replace("\\)", "").replace("\\[", "").replace("\\]", "")
    t = t.replace("R^2", "R²").replace("k = 64", "k = 64")
    t = t.replace("\\frac{1}{|M_f|}", "(1/|M_f|) ").replace("\\frac{1}{|N_f|}", "(1/|N_f|) ")
    t = t.replace("\\sum_{i \\in M_f}", "sum over members ").replace("\\sum_{j \\in N_f}", "sum over nonmembers ")
    t = t.replace("\\max_f", "max over f").replace("\\cos", "cos").replace("\\in", " in ")
    t = t.replace("\\mu_{N_f}", "mean(N_f)").replace("\\sigma_{N_f}", "SD(N_f)")
    t = t.replace("\\text{CalibratedFit}", "CalibratedFit").replace("\\text{score}", "score").replace("\\text{comments}", "comments")
    t = t.replace("\\log_{10}", "log10")
    t = re.sub(r"d_f\b", "d_f", t)
    t = t.replace("\\", "")
    return t

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
def add_runs(par, text):
    text = demath(text)
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True; r.font.name = "Times New Roman"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1]); r.italic = True; r.font.name = "Times New Roman"
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Times New Roman"
        else:
            r = par.add_run(piece); r.font.name = "Times New Roman"

def convert(md, doc):
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        # page break marker
        if ln.strip() in ("\\newpage",):
            doc.add_page_break(); i += 1; continue
        if ln.strip() == "---":
            i += 1; continue
        # tables: block of lines starting with |
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in row.strip("|").split("|")] for row in block]
            rows = [r for r in rows if not all(re.fullmatch(r"[-:\s]*", c) for c in r)]  # drop |---| sep
            if rows:
                ncol = max(len(r) for r in rows)
                tbl = doc.add_table(rows=0, cols=ncol); tbl.style = "Light Grid Accent 1"
                for ri, r in enumerate(rows):
                    cells = tbl.add_row().cells
                    for ci in range(ncol):
                        txt = r[ci] if ci < len(r) else ""
                        p = cells[ci].paragraphs[0]; add_runs(p, txt)
                        for run in p.runs:
                            run.font.name = "Times New Roman"; run.font.size = Pt(10)
                            if ri == 0: run.bold = True
            doc.add_paragraph()
            continue
        # figure caption block: insert image, then caption (strip the placeholder)
        lnf = ln.lstrip("- ").rstrip()
        figkey = next((k for k in FIGMAP if lnf.startswith(k)), None)
        if figkey:
            path = FIGMAP[figkey]
            if os.path.exists(path):
                pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(path, width=Inches(5.6))
            cap = lnf.replace(" *[Figure image to be inserted.]*", "").replace("*[Figure image to be inserted.]*", "")
            p = doc.add_paragraph(); add_runs(p, cap); i += 1
            continue
        # headings — plain 12pt Times New Roman bold (no Word heading styles)
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            p = doc.add_paragraph()
            r = p.add_run(demath(m.group(2)))
            r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0, 0, 0)
            p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            i += 1; continue
        # list item
        if re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, re.sub(r"^\s*[-*]\s+", "", ln)); i += 1; continue
        if re.match(r"^\s*\d+\.\s+", ln):
            p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\s*\d+\.\s+", "", ln)); i += 1; continue
        # blank
        if not ln.strip():
            i += 1; continue
        # normal paragraph (gather until blank)
        para = [ln]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "---")) \
                and not re.match(r"^\s*[-*]\s+", lines[i]) and not re.match(r"^\s*\d+\.\s+", lines[i]):
            para.append(lines[i].rstrip()); i += 1
        p = doc.add_paragraph(); add_runs(p, " ".join(para))

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.paragraph_format.space_after = Pt(6)
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
convert(open(MAN, encoding="utf-8").read(), doc)
doc.add_page_break()
h = doc.add_paragraph(); hr = h.add_run("SUPPLEMENTARY INFORMATION")
hr.bold = True; hr.font.name = "Times New Roman"; hr.font.size = Pt(12); hr.font.color.rgb = RGBColor(0, 0, 0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
convert(open(SUP, encoding="utf-8").read(), doc)
doc.save(OUT)
import os
print("WROTE", OUT, os.path.getsize(OUT), "bytes;", len(doc.paragraphs), "paragraphs,", len(doc.tables), "tables")
