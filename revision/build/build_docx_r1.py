#!/usr/bin/env python3
"""Revision DOCX builder (R1). Derived from the submitted builder
`SSCR Submission/_Working/07_Analysis_Audit/build_docx.py` (sha256 055b780a…; verbatim copy kept alongside as
`build_docx_submitted_copy.py`). Changes relative to the submitted builder:
  1. Inline math (R1-4 / AI-12). Each \\( ... \\) span is rendered as Word runs: `_x` / `_{..}` become subscript runs,
     `^x` / `^{..}` superscript runs, single-letter variables italic, LaTeX symbols mapped to Unicode (∈ ℝ Δ ∑ μ σ ≤ …),
     and no backslash, `_{` or `^{` is left in the text. The submitted builder stripped backslashes and left `d_f`,
     `M_f`, `z_c[T_f]` etc. as literal underscores.
  2. Display math: the w_f equation now includes its `w_f ∈ ℝ^15_+` clause (the submitted OMML dropped it).
  3. Change highlighting (ED-2): changed spans relative to the submitted _Working MS/SI are coloured (no track changes).
     Sentence-level lookup against the submitted text; an unmatched sentence is word-diffed against its closest
     submitted sentence (difflib ratio ≥ 0.6) and only the differing words are coloured, otherwise the whole sentence is.
     `--clean` builds the uncoloured version.
  4. FIGMAP points at the revision figures (19_Revision/build/out/figures); S8 withdrawn, S12 added, Figure 5 added.
  5. HTML comments (internal build notes) are dropped.
Usage (from 19_Revision/build/):
  python3 build_docx_r1.py OUT.docx {manuscript|supplement|titlepage|letter|memo|coverletter} [--clean] [--review]  (letter, memo, coverletter: --clean only; --review embeds main figures)
"""
import difflib
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Inches, Pt, RGBColor

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
REV = os.path.dirname(HERE)
SUBW = os.path.join(REV, "..", "..", "SSCR Submission", "_Working")
FIG = os.path.join(HERE, "out", "figures")
SRC = {"manuscript": os.path.join(REV, "Manuscript_R1.md"), "supplement": os.path.join(REV, "Supplement_R1.md"),
       "titlepage": os.path.join(REV, "Title_Page_R1.md"),
       # clean-only builds (09 F-68): the response letter and the ED-1 change memo have no submitted counterpart
       "letter": os.path.join(REV, "..", "21_Response_Letter.md"), "memo": os.path.join(REV, "Change_Memo_R1.md"), "coverletter": os.path.join(REV, "Cover_Letter_R1.md")}
OLD = {"manuscript": os.path.join(SUBW, "01_Manuscript", "SSCR_Manuscript_Anonymized.md"),
       "supplement": os.path.join(SUBW, "02_Supplement", "SSCR_Supplement.md"),
       "titlepage": os.path.join(SUBW, "01_Manuscript", "SSCR_Title_Page.md")}
CHANGE_RGB = RGBColor(0x1F, 0x4E, 0xB4)  # changed text: blue

PLACEHOLDERS = False
FIGMAP = {
    "**Figure 1.": "Figure1_meme_family_concept.png",
    "**Figure 2.": "Figure2_measurement_framework.png",
    "**Figure 3.": "Figure3_family_recovery.png",
    "**Figure 4.": "Figure4_adversarial.png",
    "**Figure 5.": "Figure5_decomposition.png",
    "**Figure S1.": "FigureS1_family_sizes.png",
    "**Figure S2.": "FigureS2_dedup_schematic.png",
    "**Figure S3.": "FigureS3_sensitivity.png",
    "**Figure S4.": "FigureS4_recovery_heatmap.png",
    "**Figure S5.": "FigureS5_paired_size.png",
    "**Figure S6.": "FigureS6_adversarial.png",
    "**Figure S7.": "FigureS7_worked_examples.png",
    "**Figure S9.": "FigureS9_sensitivity.png",
    "**Figure S10.": "FigureS10_decomposition_nulls.png",
    "**Figure S11.": "FigureS11_prospective.png",
    "**Figure S12.": "FigureS12_per_fold.png",
    "**Figure S13.": "FigureS13_development_diagnostics.png",
}

# ------------------------------------------------------------------ text (non-math) normalization
def demath(t):
    """Plain-text normalization for non-math text (math spans are handled by math_segs)."""
    t = t.replace("R^2", "R²").replace("\\Delta", "Δ").replace("\\times", "×")
    t = re.sub(r"Δ\s+(R|AUROC)", r"Δ\1", t)
    t = t.replace("\\", "")
    t = re.sub(r"(?<=[\s(\[=])-(?=[.\d])", "−", t)
    return t


# ------------------------------------------------------------------ inline math -> runs
SYM = [(r"\mathbb{R}", "ℝ"), (r"\qquad", "   "), (r"\quad", "  "), (r"\in", "∈"), (r"\notin", "∉"),
       (r"\Delta", "Δ"), (r"\sum", "∑"), (r"\leq", "≤"), (r"\geq", "≥"), (r"\le", "≤"), (r"\ge", "≥"),
       (r"\neq", "≠"), (r"\approx", "≈"), (r"\pm", "±"), (r"\times", "×"), (r"\cdot", "·"), (r"\ldots", "…"),
       (r"\mu", "μ"), (r"\sigma", "σ"), (r"\rho", "ρ"), (r"\alpha", "α"), (r"\beta", "β"), (r"\tau", "τ"),
       (r"\lambda", "λ"), (r"\epsilon", "ε"), (r"\left", ""), (r"\right", ""), (r"\,", " "), (r"\;", " "),
       (r"\{", "{"), (r"\}", "}"), (r"\|", "‖")]
FUNCS = ("cos", "max", "min", "log", "exp", "argmax")


def _group(s, k):
    """Return (content, next_index) for a TeX argument starting at s[k] ({..} or a single char)."""
    if k < len(s) and s[k] == "{":
        depth, j = 0, k
        while j < len(s):
            depth += s[j] == "{"
            depth -= s[j] == "}"
            if depth == 0:
                return s[k + 1:j], j + 1
            j += 1
        return s[k + 1:], len(s)
    return (s[k], k + 1) if k < len(s) else ("", k)


def _prep(tex):
    tex = re.sub(r"\\(?:text|mathrm|operatorname|textrm)\{([^{}]*)\}", r"\\TXT{\1}", tex)
    while "\\frac" in tex:
        k = tex.index("\\frac")
        a, j = _group(tex, k + 5)
        b, j = _group(tex, j)
        wrap = lambda x: x if re.fullmatch(r"[\w|μσΔ]+", x.replace("\\TXT", "")) else f"({x})"  # noqa: E731
        tex = tex[:k] + wrap(a) + "/" + wrap(b) + tex[j:]
    for f in FUNCS:
        tex = tex.replace("\\" + f, "\\TXT{" + f + "}")
    for a, b in SYM:
        tex = re.sub(re.escape(a) + r"(?![A-Za-z])", lambda _m, b=b: b, tex)
    return re.sub(r"Δ\s+", "Δ", tex)


def math_segs(tex, vert=None):
    """-> list of (text, vert, italic, upright_word) for one math span."""
    tex = _prep(tex) if vert is None else tex
    out, k, buf = [], 0, ""

    def flush():
        nonlocal buf
        for tok in re.findall(r"[A-Za-z]+|[^A-Za-z]", buf):
            out.append((tok.replace("-", "−"), vert, bool(re.fullmatch(r"[A-Za-z]", tok))))
        buf = ""
    while k < len(tex):
        ch = tex[k]
        if tex.startswith("\\TXT{", k):
            flush()
            g, k = _group(tex, k + 4)
            out.append((g, vert, False))
            continue
        if ch in "_^":
            flush()
            g, k = _group(tex, k + 1)
            out.extend(math_segs(g, "subscript" if ch == "_" else "superscript"))
            continue
        if ch in "{}":
            k += 1
            continue
        buf += ch
        k += 1
    flush()
    return out


# ------------------------------------------------------------------ paragraph segmentation
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
MATH = re.compile(r"\\\((.+?)\\\)")


def segments(text):
    """Markdown line -> [(text, fmt dict)]; fmt keys: bold, italic, vert."""
    segs = []
    for piece in INLINE.split(text):
        if not piece:
            continue
        fmt = {}
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece, fmt = piece[2:-2], {"bold": True}
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, fmt = piece[1:-1], {"italic": True}
        elif piece.startswith("`") and piece.endswith("`"):
            segs.append((piece[1:-1], {"code": True}))
            continue
        pos = 0
        for m in MATH.finditer(piece):
            if m.start() > pos:
                segs.append((demath(piece[pos:m.start()]), dict(fmt)))
            for t, v, it in math_segs(m.group(1)):
                segs.append((t, dict(fmt, vert=v, italic=fmt.get("italic") or it)))
            pos = m.end()
        if pos < len(piece):
            segs.append((demath(piece[pos:]), dict(fmt)))
    return segs


def plain(text):
    return "".join(t for t, _ in segments(text))


# ------------------------------------------------------------------ change detection against the submitted text
SENT = re.compile(r"(?<=[.!?:;])\s+|\s*\|\s*")
NORM = lambda s: re.sub(r"\s+", " ", s).strip()  # noqa: E731


class Differ:
    def __init__(self, old_md):
        self.sents = set()
        old = re.sub(r"<!--.*?-->", "", old_md, flags=re.S)
        paras, cur = [], []
        for ln in old.split("\n"):  # rebuild paragraphs as convert() does (wrapped lines are joined)
            starts_block = not ln.strip() or ln.startswith(("#", "|", "---", "\\[")) \
                or re.match(r"^\s*([-*]|\d+\.)\s+", ln)
            if starts_block and cur:
                paras.append(" ".join(cur))
                cur = []
            if ln.strip():
                cur.append(ln.strip())
                if ln.startswith(("#", "|")):
                    paras.append(" ".join(cur))
                    cur = []
        paras.append(" ".join(cur))
        for ln in paras:
            ln = re.sub(r"^\s*(#{1,4}|[-*]|\d+\.)\s+", "", ln.strip())
            if ln and not re.fullmatch(r"[|\-:\s]+", ln):
                for s in SENT.split(plain(ln)):
                    if NORM(s):
                        self.sents.add(NORM(s))
        self.pool = sorted(self.sents)
        self.n_changed = self.n_total = 0

    def mask(self, P):
        m = [False] * len(P)
        pos = 0
        for s in SENT.split(P):
            start = P.find(s, pos)
            pos = start + len(s)
            ns = NORM(s)
            if not ns:
                continue
            self.n_total += 1
            if ns in self.sents:
                continue
            self.n_changed += 1
            best = difflib.get_close_matches(ns, self.pool, n=1, cutoff=0.6)
            if not best:
                for i in range(start, start + len(s)):
                    m[i] = not s[i - start].isspace()
                continue
            # word-level diff: colour only words absent from the matched submitted sentence
            toks = [(mm.start(), mm.end(), mm.group()) for mm in re.finditer(r"\S+", s)]
            old_toks = best[0].split(" ")
            sm = difflib.SequenceMatcher(a=old_toks, b=[t for _, _, t in toks], autojunk=False)
            keep = set()
            for blk in sm.get_matching_blocks():
                keep.update(range(blk.b, blk.b + blk.size))
            for ti, (a, b, _) in enumerate(toks):
                if ti not in keep:
                    for i in range(start + a, start + b):
                        m[i] = True
        for i in range(1, len(P) - 1):  # colour a space that sits between two coloured words
            if P[i].isspace() and m[i - 1] and not m[i]:
                j = i
                while j < len(P) and P[j].isspace():
                    j += 1
                if j < len(P) and m[j]:
                    for q in range(i, j):
                        m[q] = True
        return m


DIFF = None


def add_runs(par, text, bold=False, size=None):
    segs = segments(text)
    P = "".join(t for t, _ in segs)
    mask = DIFF.mask(P) if DIFF else [False] * len(P)
    k = 0
    for t, fmt in segs:
        j = 0
        while j < len(t):
            ch = mask[k + j]
            e = j
            while e < len(t) and mask[k + e] == ch:
                e += 1
            r = par.add_run(t[j:e])
            r.font.name = "Times New Roman"
            r.bold = bold or fmt.get("bold") or None
            r.italic = fmt.get("italic") or None
            if fmt.get("vert") == "subscript":
                r.font.subscript = True
            elif fmt.get("vert") == "superscript":
                r.font.superscript = True
            if size:
                r.font.size = Pt(size)
            if ch:
                r.font.color.rgb = CHANGE_RGB
            j = e
        k += len(t)


# ------------------------------------------------------------------ OMML display equations (as submitted, plus w_f clause)
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
def _r(t): return f'<m:r><m:t xml:space="preserve">{t}</m:t></m:r>'
def _sub(b, s): return f'<m:sSub><m:e>{b}</m:e><m:sub>{s}</m:sub></m:sSub>'
def _subsup(b, s, p): return f'<m:sSubSup><m:e>{b}</m:e><m:sub>{s}</m:sub><m:sup>{p}</m:sup></m:sSubSup>'
def _frac(n, d): return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'
def _absx(x): return _r("|") + x + _r("|")
_Mf, _Nf = _sub(_r("M"), _r("f")), _sub(_r("N"), _r("f"))
EQ = {
    "d_f =": (_sub(_r("d"), _r("f")) + _r(" = ") + _frac(_r("1"), _absx(_Mf))
              + _sub(_r("∑"), _r("i ∈ ") + _Mf) + _sub(_r("z"), _r("i")) + _r(" − ")
              + _frac(_r("1"), _absx(_Nf)) + _sub(_r("∑"), _r("j ∈ ") + _Nf) + _sub(_r("z"), _r("j")) + _r(".")),
    "w_f =": (_sub(_r("w"), _r("f")) + _r(" = max(") + _sub(_r("d"), _r("f")) + _r("[")
              + _sub(_r("T"), _r("f")) + _r("], 0),    ") + _sub(_r("w"), _r("f")) + _r(" ∈ ")
              + _subsup(_r("ℝ"), _r("+"), _r("15")) + _r(".")),
    "r(c,f) =": (_r("r(c, f) = cos(") + _sub(_r("z"), _r("c")) + _r("[") + _sub(_r("T"), _r("f"))
                 + _r("], ") + _sub(_r("w"), _r("f")) + _r(").")),
    "\\max_f": _sub(_r("max"), _r("f")) + _r(" r(c, f),"),
    "CalibratedFit": (_r("CalibratedFit(c, f) = ")
                      + _frac(_r("r(c, f) − ") + _sub(_r("μ"), _Nf), _sub(_r("σ"), _Nf)) + _r(",")),
    "score}": _sub(_r("log"), _r("10")) + _r("(score + 1)."),
    "comments}": _sub(_r("log"), _r("10")) + _r("(comments + 1)."),
}


def add_equation(doc, inner):
    if doc.paragraphs:  # proof fix (F-45): the lead-in paragraph stays on the same page as its display equation
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(parse_xml(f'<m:oMath xmlns:m="{M_NS}">{inner}</m:oMath>'))
    return p


def harden_table(tbl):
    for ri, row in enumerate(tbl.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trPr.append(OxmlElement("w:tblHeader"))


def start_section(doc, landscape):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    return sec


def convert(md, doc):
    md = re.sub(r"<!--.*?-->\n?", "", md, flags=re.S)
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.strip() == "\\newpage":
            doc.add_page_break()
            i += 1
            continue
        if ln.strip() == "---":
            i += 1
            continue
        if ln.strip().startswith("\\["):
            blk, j = [], i
            while j < len(lines) and "\\]" not in lines[j]:
                blk.append(lines[j])
                j += 1
            if j < len(lines):
                blk.append(lines[j])
            tex = " ".join(x.replace("\\[", "").replace("\\]", "") for x in blk)
            key = next((k for k in EQ if k in tex), None)
            if key:
                add_equation(doc, EQ[key])
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, "\\(" + tex.strip() + "\\)")
            i = j + 1
            continue
        if ln.startswith("|"):
            prevtext = " ".join(lines[max(0, i - 4):i])
            is_t3 = "**Table 3." in prevtext
            is_s7 = "**Table S7." in prevtext
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in row.strip("|").split("|")] for row in block]
            rows = [r for r in rows if not all(re.fullmatch(r"[-:\s]*", c) for c in r)]
            if rows:
                # keep each table caption on the same page as its table (final rendering proof, `26` RS-2)
                paras = doc.paragraphs
                ci_ = next((k for k in range(len(paras) - 1, max(-1, len(paras) - 5), -1)
                            if paras[k].text.strip().startswith("Table")), None)
                cap = paras[ci_] if ci_ is not None else None
                if cap is not None:
                    for q in paras[ci_:]:
                        q.paragraph_format.keep_with_next = True
                if is_t3 and cap is not None:
                    cap.paragraph_format.page_break_before = True  # break before the caption, not between caption and table
                elif is_t3:
                    doc.add_page_break()
                if is_s7:
                    moved = []
                    if cap is not None:  # move the caption into the landscape section with its table
                        for q in paras[ci_:]:
                            moved.append(q._p)
                        for el in moved:
                            el.getparent().remove(el)
                    start_section(doc, landscape=True)
                    body = doc.element.body
                    for el in moved:
                        body.insert(len(body) - 1, el)
                ncol = max(len(r) for r in rows)
                tbl = doc.add_table(rows=0, cols=ncol)
                tbl.style = "Light Grid Accent 1"
                for ri, r in enumerate(rows):
                    cells = tbl.add_row().cells
                    for ci in range(ncol):
                        add_runs(cells[ci].paragraphs[0], r[ci] if ci < len(r) else "", bold=(ri == 0),
                                 size=9 if is_s7 else (8.5 if ncol >= 7 else 10))  # wide tables: avoid mid-word breaks
                harden_table(tbl)
                if is_s7:
                    start_section(doc, landscape=False)
            doc.add_paragraph()
            continue
        lnf = ln.lstrip("- ").rstrip()
        figkey = next((k for k in FIGMAP if lnf.startswith(k)), None)
        if figkey:
            path = os.path.join(FIG, FIGMAP[figkey])
            if not os.path.exists(path):
                raise SystemExit(f"missing figure for {figkey}: {path}")
            w = 4.9 if figkey == "**Figure S4." else 6.3
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            if PLACEHOLDERS and not figkey.startswith("**Figure S"):
                # SSCR rule (the 2026-07-21 unsubmission): main figures are not embedded in the manuscript; they are
                # uploaded as separate files, with a placeholder and the caption kept in the text, as submitted.
                pic.add_run(f"[Insert {figkey.strip('*.')} about here]")
            else:
                pic.add_run().add_picture(path, width=Inches(w))
            p = doc.add_paragraph()
            add_runs(p, lnf)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            p = doc.add_paragraph()
            add_runs(p, m.group(2), bold=True, size=12)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*[-*]\s+", "", ln))
            i += 1
            continue
        mo = re.match(r"^\s*(\d+)\.\s+(.*)", ln)
        if mo:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            add_runs(p, f"{mo.group(1)}. {mo.group(2)}")
            i += 1
            continue
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("> "):  # quoted reviewer comment (letter): indented paragraph
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            add_runs(p, ln[2:].strip())
            i += 1
            continue
        para = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "---")) \
                and not re.match(r"^\s*[-*]\s+", lines[i]) and not re.match(r"^\s*\d+\.\s+", lines[i]):
            para.append(lines[i].rstrip())
            i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(para))


def main():
    global DIFF, PLACEHOLDERS
    out, mode = sys.argv[1], sys.argv[2]
    # main figures -> "[Insert Figure N about here]" for upload; --review embeds them for the manual-review copy (09 F-69)
    PLACEHOLDERS = mode == "manuscript" and "--review" not in sys.argv
    clean = "--clean" in sys.argv
    if not clean and mode not in OLD:
        raise SystemExit(f"mode {mode!r} has no submitted counterpart: build it with --clean")
    DIFF = None if clean else Differ(open(OLD[mode], encoding="utf-8").read())
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    if not clean:
        p = doc.add_paragraph()
        r = p.add_run("Revised text relative to the submitted version is shown in blue.")
        r.italic = True
        r.font.color.rgb = CHANGE_RGB
    convert(open(SRC[mode], encoding="utf-8").read(), doc)
    doc.save(out)
    extra = f"; sentences changed {DIFF.n_changed}/{DIFF.n_total}" if DIFF else ""
    print("WROTE", out, f"(mode={mode}{', clean' if clean else ', highlighted'})", os.path.getsize(out), "bytes;",
          len(doc.paragraphs), "paragraphs,", len(doc.tables), "tables" + extra)


if __name__ == "__main__":
    main()
